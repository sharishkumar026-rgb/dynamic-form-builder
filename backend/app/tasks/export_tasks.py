from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Table,
    TableStyle,
)
from reportlab.lib.styles import getSampleStyleSheet

from app.core.config import settings
from app.core.database import SessionLocal
from app.repositories.form_repository import FormRepository
from app.repositories.response_repository import ResponseRepository
from app.tasks.celery_app import celery_app


# ============================================================
# EXPORT DIRECTORY
# ============================================================

EXPORT_DIR = Path(settings.UPLOAD_DIR) / "exports"

EXPORT_DIR.mkdir(
    parents=True,
    exist_ok=True,
)


# ============================================================
# HELPERS
# ============================================================

def _parse_date(
    value: str | None,
) -> datetime | None:
    """
    Convert an ISO datetime string into a datetime object.
    """

    if not value:
        return None

    return datetime.fromisoformat(value)


# ============================================================
# GET REPORT DATA
# ============================================================

def _get_report_data(
    db,
    form_id: int,
    start_date: datetime | None = None,
    end_date: datetime | None = None,
    submitted_by_id: int | None = None,
):
    """
    Retrieve form and response data required for exports.

    Supports filtering by:

        - form_id
        - start_date
        - end_date
        - submitted_by_id
    """

    form_repository = FormRepository(db)

    response_repository = ResponseRepository(db)

    # --------------------------------------------------------
    # Get form
    # --------------------------------------------------------

    form = form_repository.get_by_id(
        form_id
    )

    if not form:
        raise ValueError(
            "Form not found"
        )

    # --------------------------------------------------------
    # DEBUG INFORMATION
    # --------------------------------------------------------

    print("=" * 70)
    print("EXPORT DEBUG")
    print("=" * 70)

    print(
        f"Form ID       : {form_id}"
    )

    print(
        f"Start Date    : {start_date}"
    )

    print(
        f"End Date      : {end_date}"
    )

    print(
        f"Submitted By  : {submitted_by_id}"
    )

    # --------------------------------------------------------
    # Get responses
    # --------------------------------------------------------

    responses = response_repository.get_by_form(
        form_id=form_id,
        skip=0,
        limit=100000,
        submitted_by_id=submitted_by_id,
        start_date=start_date,
        end_date=end_date,
    )

    # --------------------------------------------------------
    # DEBUG RESPONSE COUNT
    # --------------------------------------------------------

    print(
        f"Responses     : {len(responses)}"
    )

    # --------------------------------------------------------
    # DEBUG RESPONSE DATA
    # --------------------------------------------------------

    for response in responses:

        print("-" * 70)

        print(
            f"Response ID   : {response.id}"
        )

        print(
            f"Form ID       : {response.form_id}"
        )

        print(
            f"Submitted At  : {response.submitted_at}"
        )

        print(
            f"Submitted By  : {response.submitted_by_id}"
        )

        print(
            f"Details       : {len(response.details)}"
        )

        for detail in response.details:

            print(
                f"  Field ID={detail.field_id} "
                f"Value={detail.value!r} "
                f"Value JSON={detail.value_json!r}"
            )

    print("=" * 70)

    return form, responses


# ============================================================
# BUILD RESPONSE ROWS
# ============================================================

def _get_response_rows(
    form,
    responses,
):
    """
    Convert response objects into rows
    suitable for Excel and PDF.
    """

    # --------------------------------------------------------
    # Sort form fields
    # --------------------------------------------------------

    fields = sorted(
        form.fields,
        key=lambda field: field.display_order,
    )

    # --------------------------------------------------------
    # Headers
    # --------------------------------------------------------

    headers = [
        "Response ID",
        "Submitted By",
        "Submitted At",
    ]

    headers.extend(
        field.label
        for field in fields
    )

    rows = []

    # --------------------------------------------------------
    # Response rows
    # --------------------------------------------------------

    for response in responses:

        submitted_by = None

        if response.submitted_by:

            submitted_by = (
                response.submitted_by.name
            )

        row = [
            response.id,
            submitted_by or "Anonymous",
            (
                response.submitted_at.strftime(
                    "%Y-%m-%d %H:%M:%S"
                )
                if response.submitted_at
                else ""
            ),
        ]

        # ----------------------------------------------------
        # Response details
        # ----------------------------------------------------

        detail_map = {
            detail.field_id: detail
            for detail in response.details
        }

        for field in fields:

            detail = detail_map.get(
                field.id
            )

            # ------------------------------------------------
            # No detail for this field
            # ------------------------------------------------

            if not detail:

                row.append("")

                continue

            # ------------------------------------------------
            # JSON value
            # ------------------------------------------------

            if detail.value_json is not None:

                if isinstance(
                    detail.value_json,
                    list,
                ):

                    row.append(
                        ", ".join(
                            map(
                                str,
                                detail.value_json,
                            )
                        )
                    )

                elif isinstance(
                    detail.value_json,
                    dict,
                ):

                    row.append(
                        str(
                            detail.value_json
                        )
                    )

                else:

                    row.append(
                        str(
                            detail.value_json
                        )
                    )

            # ------------------------------------------------
            # Normal text value
            # ------------------------------------------------

            else:

                row.append(
                    detail.value or ""
                )

        rows.append(row)

    # --------------------------------------------------------
    # DEBUG ROW INFORMATION
    # --------------------------------------------------------

    print("=" * 70)
    print("EXPORT ROW DEBUG")
    print("=" * 70)

    print(
        f"Headers : {headers}"
    )

    print(
        f"Rows    : {len(rows)}"
    )

    for row in rows:

        print(
            f"Row     : {row}"
        )

    print("=" * 70)

    return headers, rows


# ============================================================
# EXCEL EXPORT
# ============================================================

@celery_app.task(
    bind=True,
    name="app.tasks.export_tasks.generate_excel_report",
)
def generate_excel_report(
    self,
    form_id: int,
    start_date: str | None = None,
    end_date: str | None = None,
    submitted_by_id: int | None = None,
    user_id: int | None = None,
):
    """
    Generate an Excel report for a form.

    Supports optional filtering by:

        - start_date
        - end_date
        - submitted_by_id

    user_id represents the user who requested
    the export.
    """

    db = SessionLocal()

    try:

        # ----------------------------------------------------
        # Parse dates
        # ----------------------------------------------------

        parsed_start_date = _parse_date(
            start_date
        )

        parsed_end_date = _parse_date(
            end_date
        )

        # ----------------------------------------------------
        # Get form and responses
        # ----------------------------------------------------

        form, responses = _get_report_data(
            db=db,
            form_id=form_id,
            start_date=parsed_start_date,
            end_date=parsed_end_date,
            submitted_by_id=submitted_by_id,
        )

        # ----------------------------------------------------
        # Build response rows
        # ----------------------------------------------------

        headers, rows = _get_response_rows(
            form=form,
            responses=responses,
        )

        # ----------------------------------------------------
        # Create workbook
        # ----------------------------------------------------

        workbook = Workbook()

        worksheet = workbook.active

        worksheet.title = "Responses"

        # ----------------------------------------------------
        # Form information
        # ----------------------------------------------------

        worksheet["A1"] = "Form Report"

        worksheet["A2"] = "Form ID"
        worksheet["B2"] = form.id

        worksheet["A3"] = "Form Title"
        worksheet["B3"] = form.title

        worksheet["A4"] = "Total Responses"
        worksheet["B4"] = len(
            responses
        )

        if submitted_by_id is not None:

            worksheet["A5"] = (
                "Submitted By User ID"
            )

            worksheet["B5"] = (
                submitted_by_id
            )

        # ----------------------------------------------------
        # Header row
        # ----------------------------------------------------

        header_row = 7

        for column_index, header in enumerate(
            headers,
            start=1,
        ):

            worksheet.cell(
                row=header_row,
                column=column_index,
                value=header,
            )

        # ----------------------------------------------------
        # Data rows
        # ----------------------------------------------------

        for row_index, row in enumerate(
            rows,
            start=header_row + 1,
        ):

            for column_index, value in enumerate(
                row,
                start=1,
            ):

                worksheet.cell(
                    row=row_index,
                    column=column_index,
                    value=value,
                )

        # ----------------------------------------------------
        # Column widths
        # ----------------------------------------------------

        for column in worksheet.columns:

            max_length = 0

            for cell in column:

                value = str(
                    cell.value or ""
                )

                if len(value) > max_length:

                    max_length = len(value)

            worksheet.column_dimensions[
                column[0].column_letter
            ].width = min(
                max_length + 2,
                40,
            )

        # ----------------------------------------------------
        # File name
        # ----------------------------------------------------

        timestamp = datetime.now().strftime(
            "%Y%m%d_%H%M%S"
        )

        filename = (
            f"form_{form_id}_report_"
            f"{timestamp}.xlsx"
        )

        file_path = (
            EXPORT_DIR / filename
        )

        # ----------------------------------------------------
        # Save workbook
        # ----------------------------------------------------

        workbook.save(
            file_path
        )

        # ----------------------------------------------------
        # Return result
        # ----------------------------------------------------

        return {
            "success": True,
            "task_id": self.request.id,
            "form_id": form_id,
            "file_name": filename,
            "file_path": str(file_path),
            "file_type": "xlsx",
            "total_responses": len(
                responses
            ),
            "submitted_by_id": submitted_by_id,
            "requested_by": user_id,
        }

    finally:

        db.close()


# ============================================================
# PDF EXPORT
# ============================================================

@celery_app.task(
    bind=True,
    name="app.tasks.export_tasks.generate_pdf_report",
)
def generate_pdf_report(
    self,
    form_id: int,
    start_date: str | None = None,
    end_date: str | None = None,
    submitted_by_id: int | None = None,
    user_id: int | None = None,
):
    """
    Generate a PDF report for a form.

    Supports optional filtering by:

        - start_date
        - end_date
        - submitted_by_id

    user_id represents the user who requested
    the export.
    """

    db = SessionLocal()

    try:

        # ----------------------------------------------------
        # Parse dates
        # ----------------------------------------------------

        parsed_start_date = _parse_date(
            start_date
        )

        parsed_end_date = _parse_date(
            end_date
        )

        # ----------------------------------------------------
        # Get form and responses
        # ----------------------------------------------------

        form, responses = _get_report_data(
            db=db,
            form_id=form_id,
            start_date=parsed_start_date,
            end_date=parsed_end_date,
            submitted_by_id=submitted_by_id,
        )

        # ----------------------------------------------------
        # Build response rows
        # ----------------------------------------------------

        headers, rows = _get_response_rows(
            form=form,
            responses=responses,
        )

        # ----------------------------------------------------
        # File name
        # ----------------------------------------------------

        timestamp = datetime.now().strftime(
            "%Y%m%d_%H%M%S"
        )

        filename = (
            f"form_{form_id}_report_"
            f"{timestamp}.pdf"
        )

        file_path = (
            EXPORT_DIR / filename
        )

        # ----------------------------------------------------
        # Create PDF document
        # ----------------------------------------------------

        document = SimpleDocTemplate(
            str(file_path),
            pagesize=landscape(A4),
            rightMargin=20,
            leftMargin=20,
            topMargin=20,
            bottomMargin=20,
        )

        styles = getSampleStyleSheet()

        elements = []

        # ----------------------------------------------------
        # Title
        # ----------------------------------------------------

        elements.append(
            Paragraph(
                f"Form Report: {form.title}",
                styles["Title"],
            )
        )

        # ----------------------------------------------------
        # Form ID
        # ----------------------------------------------------

        elements.append(
            Paragraph(
                f"Form ID: {form.id}",
                styles["Normal"],
            )
        )

        # ----------------------------------------------------
        # Total responses
        # ----------------------------------------------------

        elements.append(
            Paragraph(
                f"Total Responses: "
                f"{len(responses)}",
                styles["Normal"],
            )
        )

        # ----------------------------------------------------
        # Submitted user filter
        # ----------------------------------------------------

        if submitted_by_id is not None:

            elements.append(
                Paragraph(
                    f"Submitted By User ID: "
                    f"{submitted_by_id}",
                    styles["Normal"],
                )
            )

        # ----------------------------------------------------
        # Space
        # ----------------------------------------------------

        elements.append(
            Paragraph(
                "<br/>",
                styles["Normal"],
            )
        )

        # ----------------------------------------------------
        # Table data
        # ----------------------------------------------------

        table_data = [
            headers
        ]

        for row in rows:

            table_data.append(
                [
                    str(
                        value if value is not None else ""
                    )
                    for value in row
                ]
            )

        # ----------------------------------------------------
        # Create table
        # ----------------------------------------------------

        table = Table(
            table_data,
            repeatRows=1,
        )

        # ----------------------------------------------------
        # Table style
        # ----------------------------------------------------

        table.setStyle(
            TableStyle(
                [
                    (
                        "BACKGROUND",
                        (0, 0),
                        (-1, 0),
                        colors.grey,
                    ),
                    (
                        "TEXTCOLOR",
                        (0, 0),
                        (-1, 0),
                        colors.white,
                    ),
                    (
                        "FONTNAME",
                        (0, 0),
                        (-1, 0),
                        "Helvetica-Bold",
                    ),
                    (
                        "FONTSIZE",
                        (0, 0),
                        (-1, -1),
                        7,
                    ),
                    (
                        "GRID",
                        (0, 0),
                        (-1, -1),
                        0.5,
                        colors.black,
                    ),
                    (
                        "VALIGN",
                        (0, 0),
                        (-1, -1),
                        "TOP",
                    ),
                    (
                        "ROWBACKGROUNDS",
                        (0, 1),
                        (-1, -1),
                        [
                            colors.white,
                            colors.lightgrey,
                        ],
                    ),
                ]
            )
        )

        elements.append(
            table
        )

        # ----------------------------------------------------
        # Build PDF
        # ----------------------------------------------------

        document.build(
            elements
        )

        # ----------------------------------------------------
        # Return result
        # ----------------------------------------------------

        return {
            "success": True,
            "task_id": self.request.id,
            "form_id": form_id,
            "file_name": filename,
            "file_path": str(file_path),
            "file_type": "pdf",
            "total_responses": len(
                responses
            ),
            "submitted_by_id": submitted_by_id,
            "requested_by": user_id,
        }

    finally:

        db.close()